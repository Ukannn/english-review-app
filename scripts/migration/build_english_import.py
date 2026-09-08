#!/usr/bin/env python3
"""Build staged and promoted SQL files from a private Sheet snapshot."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any, Iterable
from functools import partial
from snapshot import load_snapshot
from export_xlsx_snapshot import canonical_hash, write_json
from validation import validate_source, validate_targets, sql_assertions, sql_preflight
from zoneinfo import ZoneInfo
from openpyxl.utils.datetime import from_excel

ID_NAMESPACE = uuid.UUID("358cb5d4-cac7-4cbf-938b-536a13fe299b")
SHANGHAI = ZoneInfo("Asia/Shanghai")


def stable_id(kind: str, key: Any, *, owner: str) -> str:
    return str(uuid.uuid5(ID_NAMESPACE, f"{owner}:{kind}:{key}"))


def is_blank(value: Any) -> bool:
    return value is None or value == ""


def text(value: Any) -> str | None:
    return None if is_blank(value) else str(value)


def integer(value: Any, default: int | None = None) -> int | None:
    if is_blank(value):
        return default
    parsed = float(value)
    if not parsed.is_integer():
        raise ValueError("Fractional value in integer column")
    return int(parsed)


def number(value: Any) -> str | None:
    if is_blank(value):
        return None
    return str(value)


def source_row_number(value: Any) -> int | None:
    if is_blank(value):
        return None
    match = re.search(r"(\d+)$", str(value))
    return int(match.group(1)) if match else None


def boolean(value: Any) -> bool | None:
    if is_blank(value):
        return None
    normalized = str(value).strip().lower()
    if normalized in {"yes", "true", "1", "y", "resolved"}:
        return True
    if normalized in {"no", "false", "0", "n", "unresolved"}:
        return False
    return None


def json_data(value: Any, default: Any) -> Any:
    if is_blank(value):
        return default
    if isinstance(value, (dict, list, int, float, bool)):
        return value
    try:
        return json.loads(str(value))
    except json.JSONDecodeError:
        return value


def date_value(value: Any) -> str | None:
    if is_blank(value):
        return None
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.date().isoformat()
        if isinstance(converted, date):
            return converted.isoformat()
        raise ValueError("Time-only value in date column")
    raw = str(value).strip()
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        return (parsed.astimezone(SHANGHAI) if parsed.tzinfo else parsed).date().isoformat()
    except ValueError as exc:
        raise ValueError("Invalid date in source; do not guess a replacement") from exc


def timestamp_value(value: Any, base_date: str | None = None) -> str | None:
    if is_blank(value):
        return None
    # Explicit legacy missing-data marker; the original marker remains in import_rows.
    if isinstance(value, str) and value.strip().lower() == "not captured":
        return None
    if isinstance(value, (int, float)):
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.replace(tzinfo=SHANGHAI).astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
        if isinstance(converted, time) and base_date is not None:
            parsed = datetime.combine(date.fromisoformat(base_date), converted).replace(tzinfo=SHANGHAI)
            return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
        if isinstance(converted, time):
            raise ValueError("Time-only value requires its source learning date")
        if isinstance(converted, date):
            parsed = datetime.combine(converted, time()).replace(tzinfo=SHANGHAI)
            return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
    raw = str(value).strip()
    try:
        if len(raw) <= 8 and ":" in raw:
            if base_date is None:
                raise ValueError("Time-only value requires its source learning date")
            parsed = datetime.combine(date.fromisoformat(base_date), time.fromisoformat(raw))
        else:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=SHANGHAI)
        return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
    except ValueError as exc:
        raise ValueError("Invalid timestamp in source; do not substitute import time") from exc


def csv_cell(value: Any) -> str:
    if value is None:
        return r"\N"
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


class SQLBuffer(io.StringIO):
    def __init__(self):
        super().__init__()
        self.tables = {}
        self.copy_blocks = []
        self.table_columns = {}


def json_columns_for(table):
    columns = {"raw_json", "formula_cells", "selected_spans", "expected_answers", "accepted_variants",
               "candidate_suggestions", "extra_practice", "confirmation"}
    if table == "english_private.commit_journal":
        columns.add("result")
    if table == "english_private.settings":
        columns.add("value")
    return columns


def select_columns(table, columns):
    return ",".join(f"coalesce({column},'null'::jsonb)" if column in json_columns_for(table) else column for column in columns)


def emit_copy(handle, table: str, columns: list[str], rows: Iterable[Iterable[Any]]) -> None:
    records = [dict(zip(columns, row, strict=True)) for row in rows]
    json_columns = json_columns_for(table)
    if table == "english_private.settings":
        json_columns.add("value")
        for record in records:
            record["value"] = json.loads(record["value"])
    handle.tables[table] = records
    handle.table_columns[table] = columns
    start = handle.tell()
    handle.write(f"copy {table} ({','.join(columns)}) from stdin with (format csv, null '\\N');\n")
    for record in records:
        # Quote every non-null cell: a literal \N must never silently become SQL NULL.
        cells = []
        for column, value in record.items():
            if value is None and column not in json_columns:
                cells.append(r"\N")
            else:
                cell = (json.dumps(value, ensure_ascii=False, separators=(",", ":")) if column in json_columns else csv_cell(value)).replace('"', '""')
                cells.append('"' + cell + '"')
        handle.write(','.join(cells) + '\n')
    handle.write("\\.\n")
    api = ""
    for offset in range(0, len(records), 100):
        payload = sql_literal(json.dumps(records[offset:offset+100], ensure_ascii=False, separators=(",", ":")))
        api += f"insert into {table} ({','.join(columns)}) select {select_columns(table, columns)} from jsonb_populate_recordset(null::{table},{payload}::jsonb);\n"
    handle.copy_blocks.append((handle.getvalue()[start:], api))


def build_import(snapshot: Path, output: Path, owner_id: str, *, imported_at: str | None = None) -> dict:
    owner = str(uuid.UUID(str(owner_id)))
    if uuid.UUID(owner).int in (0, 1):
        raise ValueError("Use the actual Auth owner, not the old rehearsal owner")
    if output.exists() and any(output.iterdir()):
        raise ValueError("Import output must be empty; never overwrite existing evidence")
    stable_id = partial(globals()["stable_id"], owner=owner)
    manifest, sheets, formulas = load_snapshot(snapshot)
    source_audit = validate_source(sheets)
    import_timestamp = timestamp_value(imported_at or manifest.get("exported_at") or datetime.now(timezone.utc).isoformat())
    batch_id = stable_id("import-batch", manifest["snapshot_sha256"])

    contexts = {text(row["Context ID"]): stable_id("context", row["Context ID"]) for row in sheets["Context Inbox"]}
    candidates = {text(row["Candidate ID"]): stable_id("candidate", row["Candidate ID"]) for row in sheets["Candidate Bank"]}
    phrases = {text(row["ID"]): stable_id("phrase", row["ID"]) for row in sheets["Phrase Bank"]}
    queues = {text(row["Queue ID"]): stable_id("queue", row["Queue ID"]) for row in sheets["Daily Queue"]}

    session_to_queue: dict[str, str] = {}
    for sheet in ("Daily Queue", "Session Questions", "Answer Drafts", "Commit Journal"):
        for row in sheets[sheet]:
            session_key = text(row.get("Session ID"))
            queue_key = text(row.get("Queue ID"))
            if session_key and queue_key:
                previous = session_to_queue.setdefault(session_key, queue_key)
                if previous != queue_key:
                    raise ValueError(f"Session {session_key} maps to two queues")
    session_records = list(sheets["Session Log"])
    logged_sessions = {text(row["Session ID"]) for row in session_records}
    for session_key, queue_key in session_to_queue.items():
        if session_key in logged_sessions:
            continue
        queue_rows = [r for r in sheets["Daily Queue"] if text(r["Queue ID"]) == queue_key]
        if not queue_rows:
            raise ValueError("Unlogged historical session has no source queue/date")
        first = queue_rows[0]
        session_records.append({"Session ID": session_key, "Date": first["Queue Date"],
            "Max Questions": integer(first["Planned Count"], len(queue_rows)), "Scheduled Start": None,
            "Actual Start": None, "Contract Version": first["Contract Version"]})
    for row in session_records:
        session_key = text(row["Session ID"])
        if session_key not in session_to_queue:
            synthetic = f"legacy-session:{session_key}"
            session_to_queue[session_key] = synthetic
            queues[synthetic] = stable_id("queue", synthetic)

    sessions = {text(row["Session ID"]): stable_id("session", row["Session ID"]) for row in session_records}
    submissions = {text(row["Submission ID"]): stable_id("submission", row["Submission ID"]) for row in sheets["Commit Journal"]}

    queue_items = {
        (text(row["Queue ID"]), integer(row["Position"])): stable_id("queue-item", f"{row['Queue ID']}:{integer(row['Position'])}")
        for row in sheets["Daily Queue"]
    }
    active_questions = {}
    for row in sheets["Session Questions"]:
        if text(row["Question Status"]) != "rejected":
            active_questions[(text(row["Queue ID"]), integer(row["Position"]))] = stable_id("question", f"{row['Queue ID']}:{integer(row['Position'])}")

    session_questions = {}
    for row in sheets["Session Questions"]:
        if text(row["Question Status"]) != "rejected" and text(row["Session ID"]):
            session_questions[(text(row["Session ID"]), integer(row["Position"]))] = active_questions[(text(row["Queue ID"]), integer(row["Position"]))]

    draft_ids = {
        (text(row["Session ID"]), integer(row["Position"])): stable_id("answer-draft", f"{row['Session ID']}:{integer(row['Position'])}")
        for row in sheets["Answer Drafts"]
    }
    request_ids = {
        (text(row["Submission ID"]), integer(row["Position"])): stable_id("grade-request", f"{row['Submission ID']}:{integer(row['Position'])}")
        for row in sheets["Grade Requests"]
    }

    target_by_sheet = {
        "README": ("legacy_only", lambda row: None),
        "Candidate Generation Inbox": ("candidate_generation_rows", lambda row: stable_id("candidate-generation", f"{row.get('Request ID')}:{row.get('Position')}")),
        "Session Questions": ("question_attempts", lambda row: stable_id("question-attempt", f"{row['__source_row']}:{row['__row_hash']}")),
        "Answer Drafts": ("answer_drafts", lambda row: draft_ids[(text(row["Session ID"]), integer(row["Position"]))]),
        "Grade Inbox": ("grade_result_attempts", lambda row: stable_id("grade-attempt", f"{row['__source_row']}:{row['__row_hash']}")),
        "Commit Journal": ("commit_journal", lambda row: stable_id("commit-journal", row["Submission ID"])),
        "Phrase Bank": ("phrases", lambda row: phrases[text(row["ID"])]),
        "Review Log": ("review_events", lambda row: stable_id("review-event", f"{row['__source_row']}:{row['__row_hash']}")),
        "Error Log": ("error_events", lambda row: stable_id("error-event", row.get("Error ID") or row["__row_hash"])),
        "Source Notes": ("source_notes", lambda row: stable_id("source-note", row["__source_row"])),
        "Config": ("settings", lambda row: stable_id("setting", row["Key"])),
        "Candidate Bank": ("candidates", lambda row: candidates[text(row["Candidate ID"])]),
        "Daily Queue": ("daily_queue_items", lambda row: queue_items[(text(row["Queue ID"]), integer(row["Position"]))]),
        "Session Log": ("sessions", lambda row: sessions[text(row["Session ID"])]),
        "Context Inbox": ("contexts", lambda row: contexts[text(row["Context ID"])]),
        "Context Candidate Inbox": ("context_candidates", lambda row: stable_id("context-candidate", f"{row['Context ID']}:{row['Proposal Position']}")),
        "Answer Draft History": ("answer_draft_history", lambda row: stable_id("draft-history", row["History ID"])),
        "Grade Requests": ("grade_requests", lambda row: request_ids[(text(row["Submission ID"]), integer(row["Position"]))]),
    }

    stage_path = output / "01-stage.sql"
    stage = SQLBuffer()
    if True:
        out = stage
        out.write("begin;\n")
        out.write(f"select set_config('request.jwt.claim.sub','{owner}',true);\n")
        out.write(
            "insert into english_private.import_batches(id,owner_id,source_file,source_sha256,snapshot_sha256,status,expected_sheet_count,expected_row_count,expected_formula_count) values "
            f"('{batch_id}','{owner}',{sql_literal(manifest['source'])},'{manifest['source_sha256']}','{manifest['snapshot_sha256']}','staged',{manifest['sheet_count']},{sum(s['row_count'] for s in manifest['sheets'])},{manifest['formula_count']});\n"
        )
        import_rows = []
        for sheet_summary in manifest["sheets"]:
            sheet = sheet_summary["sheet"]
            target_table, target_fn = target_by_sheet[sheet]
            for row in sheets[sheet]:
                raw = {key: value for key, value in row.items() if not key.startswith("__")}
                raw["_source_headers"] = row["__headers"]
                raw["_source_values"] = row["__values"]
                target_id = target_fn(row)
                legacy_id = text(raw.get("ID")) or next((text(raw.get(key)) for key in raw if key.endswith(" ID") and text(raw.get(key))), None)
                import_rows.append((
                    stable_id("import-row", f"{batch_id}:{sheet}:{row['__source_row']}"), owner, batch_id, sheet,
                    row["__source_row"], legacy_id, raw, row["__row_hash"], formulas.get((sheet, row["__source_row"]), []),
                    target_table, target_id, "legacy_only" if target_table == "legacy_only" else "staged"
                ))
        emit_copy(out, "english_private.import_rows", [
            "id","owner_id","import_batch_id","source_sheet","source_row","legacy_id","raw_json","row_hash",
            "formula_cells","target_table","target_id","promotion_status"
        ], import_rows)
        out.write("commit;\n")

    promote_path = output / "02-promote.sql"
    promote = SQLBuffer()
    if True:
        out = promote
        out.write("begin;\n")
        out.write(f"select set_config('request.jwt.claim.sub','{owner}',true);\n")

        emit_copy(out, "english_private.settings", ["id","owner_id","key","value","revision","created_at","updated_at"], (
            (stable_id("setting", row["Key"]),owner,text(row["Key"]),json.dumps(text(row["Value"]),ensure_ascii=False),1,import_timestamp,import_timestamp)
            for row in sheets["Config"]
        ))
        emit_copy(out, "english_private.source_notes", ["id","owner_id","legacy_row","date_added","source","context","candidate_chunk","why_useful","added_to_bank","created_at"], (
            (stable_id("source-note", row["__source_row"]),owner,row["__source_row"],timestamp_value(row["Date Added"]),text(row["Source"]),text(row["Context"]),text(row["Candidate Chunk"]),text(row["Why Useful"]),boolean(row["Added to Bank?"]),timestamp_value(row["Date Added"]) or import_timestamp)
            for row in sheets["Source Notes"]
        ))
        context_status = {"pending":"pending","processed":"completed","explanation_only":"legacy"}
        emit_copy(out, "english_private.contexts", ["id","owner_id","legacy_context_id","raw_text","selected_spans","source_url","source_title","user_note","status","processing_batch_id","capture_request_id","contract_version","created_at","processed_at"], (
            (contexts[text(row["Context ID"])],owner,text(row["Context ID"]),text(row["Raw Text"]) or "[UNKNOWN]",json_data(row["Selected Spans JSON"],[]),text(row["Source URL"]),text(row["Source Title"]),text(row["User Note"]),context_status.get(text(row["Processing Status"]),"legacy"),text(row["Processing Batch ID"]),text(row["Capture Request ID"]),text(row["Contract Version"]),timestamp_value(row["Created At"]) or import_timestamp,timestamp_value(row["Processed At"]))
            for row in sheets["Context Inbox"]
        ))
        emit_copy(out, "english_private.candidates", ["id","owner_id","legacy_candidate_id","candidate","candidate_type","cue_zh","source","source_context","why_useful","topic","difficulty","natural_example","common_mistake","origin_type","origin_context_id","selected_text","source_url","intake_priority","status","deferred_reason","legacy_source_note_row","created_at","promoted_at"], (
            (candidates[text(row["Candidate ID"])],owner,text(row["Candidate ID"]),text(row["Candidate"]) or "[UNKNOWN]",text(row["Candidate Type"]),text(row["中文提示"]),text(row["Source"]),text(row["Context"]),text(row["Why Useful"]),text(row["Topic"]),text(row["Difficulty"]),text(row["Natural Example"]),text(row["Common Mistake"]),text(row["Origin Type"]),contexts.get(text(row["Origin Context ID"])),text(row["Selected Text"]),text(row["Source URL"]),text(row["Intake Priority"]),text(row["Status"]) or "ready",text(row["Deferral Reason"]),source_row_number(row["Source Note Row"]),timestamp_value(row["Date Added"]) or import_timestamp,timestamp_value(row["Date Promoted"]))
            for row in sheets["Candidate Bank"]
        ))
        generation_columns = ["id", "owner_id", "request_id", "queue_date", "requested_count", "available_count", "shortfall_count", "position", "candidate", "cue_zh", "candidate_type", "source", "source_context", "why_useful", "topic", "difficulty", "natural_example", "common_mistake", "generation_batch_id", "model_id", "status", "candidate_id", "contract_version", "created_at", "committed_at"]
        generation_rows = []
        for row in sheets["Candidate Generation Inbox"]:
            generation_rows.append((stable_id("candidate-generation", f"{row['Request ID']}:{row['Position']}"), owner,
                text(row["Request ID"]), date_value(row["Queue Date"]), integer(row["Requested Count"]), integer(row["Available Count"]),
                integer(row["Shortfall Count"]), integer(row["Position"]), text(row["Candidate"]) or "[UNKNOWN]",
                text(row["Chinese Cue"]), text(row["Candidate Type"]), text(row["Source"]), text(row["Context"]), text(row["Why Useful"]),
                text(row["Topic"]), text(row["Difficulty"]), text(row["Natural Example"]), text(row["Common Mistake"]),
                text(row["Generation Batch ID"]), text(row["Model ID"]), text(row["Generation Status"]) or "legacy",
                candidates.get(text(row["Candidate ID"])), text(row["Contract Version"]), timestamp_value(row["Created At"]) or import_timestamp,
                timestamp_value(row["Committed At"])))
        emit_copy(out, "english_private.candidate_generation_rows", generation_columns, generation_rows)
        emit_copy(out, "english_private.phrases", ["id","owner_id","legacy_phrase_id","chunk","cue_zh","phrase_type","topic","difficulty","status","review_stage","next_review_at","common_mistake","natural_example","notes","source","source_candidate_id","mastery_streak","last_result","contract_version","canonical_pattern","created_at"], (
            (phrases[text(row["ID"])],owner,text(row["ID"]),text(row["Chunk"]) or "[UNKNOWN]",text(row["中文提示"]),text(row["Type"]),text(row["Topic"]),text(row["Difficulty"]),text(row["Status"]) or "active",integer(row["Review Stage"],0),timestamp_value(row["Next Review"]),text(row["Common Mistake"]),text(row["Natural Example"]),text(row["Notes"]),text(row["Source"]),candidates.get(text(row["Source Candidate ID"])),integer(row["Mastery Streak"],0),text(row["Last Result"]),text(row["Contract Version"]),text(row["Canonical Pattern"]),timestamp_value(row["Created Date"]) or import_timestamp)
            for row in sheets["Phrase Bank"]
        ))
        promoted_pairs = [(phrases[text(row["Promoted Phrase ID"])], candidates[text(row["Candidate ID"])]) for row in sheets["Candidate Bank"] if text(row["Promoted Phrase ID"])]
        for phrase_id, candidate_id in promoted_pairs:
            out.write(f"update english_private.candidates set promoted_phrase_id='{phrase_id}' where id='{candidate_id}';\n")

        queue_groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for row in sheets["Daily Queue"]:
            queue_groups[text(row["Queue ID"])].append(row)
        for session_key, queue_key in session_to_queue.items():
            if queue_key.startswith("legacy-session:"):
                queue_groups[queue_key] = []
        queue_rows = []
        session_log_by_id = {text(row["Session ID"]): row for row in session_records}
        status_priority = {"superseded":0,"planned":1,"deferred":2,"presented":3,"committed":4}
        for queue_key, rows in sorted(queue_groups.items()):
            if rows:
                first = rows[0]
                status = max((text(row["Queue Status"]) or "planned" for row in rows), key=lambda value: status_priority.get(value,0))
                queue_rows.append((queues[queue_key],owner,queue_key,date_value(first["Queue Date"]),"legacy",integer(first["Planned Count"],len(rows)),integer(first["Adjusted Target"]),text(first["Queue Kind"]),integer(first["Plan Revision"],1),queues.get(text(first["Superseded By"])),timestamp_value(first["Superseded At"]),text(first["Change Reason"]),text(first["Contract Version"]),timestamp_value(first["Created At"]) or import_timestamp,timestamp_value(first["Committed At"])))
            else:
                session_key = queue_key.split(":",1)[1]
                session_row = session_log_by_id[session_key]
                learning_date = date_value(session_row["Date"]) or session_key[:10]
                queue_rows.append((queues[queue_key],owner,queue_key,learning_date,"legacy",integer(session_row["Max Questions"],20),integer(session_row["Max Questions"],20),"legacy",1,None,None,"Synthesized for pre-queue session",text(session_row["Contract Version"]),timestamp_value(session_row["Date"]) or import_timestamp,timestamp_value(session_row["Date"])))
        emit_copy(out, "english_private.daily_queues", ["id","owner_id","legacy_queue_id","queue_date","status","planned_count","adjusted_target","queue_kind","revision","superseded_by","superseded_at","change_reason","contract_version","created_at","committed_at"], queue_rows)
        emit_copy(out, "english_private.daily_queue_items", ["id","owner_id","queue_id","position","selection_type","phrase_id","candidate_id","chunk","cue_zh","topic","difficulty","natural_example","original_next_review","priority_reason","status","presented_at","contract_version","created_at"], (
            (queue_items[(text(row["Queue ID"]),integer(row["Position"]))],owner,queues[text(row["Queue ID"])],integer(row["Position"]),text(row["Selection Type"]),phrases.get(text(row["Phrase ID"])),candidates.get(text(row["Candidate ID"])),text(row["Chunk"]),text(row["中文提示"]),text(row["Topic"]),text(row["Difficulty"]),text(row["Natural Example"]),timestamp_value(row["Original Next Review"]),text(row["Priority Reason"]),text(row["Queue Status"]) or "planned",timestamp_value(row["Presented At"]),text(row["Contract Version"]),timestamp_value(row["Created At"]) or import_timestamp)
            for row in sheets["Daily Queue"]
        ))

        draft_sessions = {text(row["Session ID"]) for row in sheets["Answer Drafts"] if text(row["Submit Status"]) == "draft"}
        emit_copy(out, "english_private.sessions", ["id","owner_id","legacy_session_id","queue_id","learning_date","status","max_questions","revision","scheduled_start","started_at","submitted_at","completed_at","contract_version","created_at"], (
            (sessions[text(row["Session ID"])],owner,text(row["Session ID"]),queues[session_to_queue[text(row["Session ID"])]],date_value(row["Date"]) or text(row["Session ID"])[:10],"legacy_recovery",integer(row["Max Questions"],20),1,timestamp_value(row["Scheduled Start"],date_value(row["Date"])),timestamp_value(row["Actual Start"],date_value(row["Date"])),None,None,text(row["Contract Version"]),timestamp_value(row["Date"]) or import_timestamp)
            for row in session_records
        ))

        emit_copy(out, "english_private.question_attempts", ["id","owner_id","queue_id","queue_item_id","session_id","position","phrase_id","candidate_id","question_type","prompt_zh","prompt_en","expected_answers","accepted_variants","semantic_boundary","grading_rubric","generation_id","model_id","prompt_version","content_hash","question_status","legacy_session_id","contract_version","created_at","bound_at"], (
            (stable_id("question-attempt",f"{row['__source_row']}:{row['__row_hash']}"),owner,queues[text(row["Queue ID"])],queue_items.get((text(row["Queue ID"]),integer(row["Position"]))),sessions.get(text(row["Session ID"])),integer(row["Position"]),phrases.get(text(row["Phrase ID"])),candidates.get(text(row["Candidate ID"])),text(row["Question Type"]),text(row["Prompt ZH"]),text(row["Prompt EN"]),json_data(row["Expected Answers JSON"],[]),json_data(row["Accepted Variants JSON"],[]),text(row["Semantic Boundary"]),text(row["Grading Rubric"]),text(row["Generation ID"]),text(row["Model ID"]),text(row["Prompt Version"]),text(row["Content Hash"]),text(row["Question Status"]) or "legacy",text(row["Session ID"]),text(row["Contract Version"]),timestamp_value(row["Created At"]),timestamp_value(row["Bound At"]))
            for row in sheets["Session Questions"]
        ))
        emit_copy(out, "english_private.questions", ["id","owner_id","queue_id","queue_item_id","session_id","position","phrase_id","candidate_id","question_type","prompt_zh","prompt_en","expected_answers","accepted_variants","semantic_boundary","grading_rubric","generation_id","model_id","prompt_version","content_hash","status","contract_version","created_at","bound_at"], (
            (active_questions[(text(row["Queue ID"]),integer(row["Position"]))],owner,queues[text(row["Queue ID"])],queue_items.get((text(row["Queue ID"]),integer(row["Position"]))),sessions.get(text(row["Session ID"])),integer(row["Position"]),phrases.get(text(row["Phrase ID"])),candidates.get(text(row["Candidate ID"])),text(row["Question Type"]),text(row["Prompt ZH"]),text(row["Prompt EN"]),json_data(row["Expected Answers JSON"],[]),json_data(row["Accepted Variants JSON"],[]),text(row["Semantic Boundary"]),text(row["Grading Rubric"]),text(row["Generation ID"]),text(row["Model ID"]),text(row["Prompt Version"]),text(row["Content Hash"]) or hashlib.sha256(f"{row['Queue ID']}:{row['Position']}".encode()).hexdigest(),"legacy",text(row["Contract Version"]),timestamp_value(row["Created At"]) or import_timestamp,timestamp_value(row["Bound At"]))
            for row in sheets["Session Questions"] if text(row["Question Status"]) != "rejected"
        ))

        emit_copy(out, "english_private.answer_drafts", ["id","owner_id","session_id","question_id","position","answer","revision","reveal_hash","answer_hash","submit_status","idempotency_key","contract_version","created_at","updated_at"], (
            (draft_ids[(text(row["Session ID"]),integer(row["Position"]))],owner,sessions[text(row["Session ID"])],session_questions[(text(row["Session ID"]),integer(row["Position"]))],integer(row["Position"]),text(row["Answer"]) or "[UNKNOWN]",integer(row["Revision"],1),text(row["Answer Hash"]) or hashlib.sha256(f"reveal:{row['Session ID']}:{row['Position']}".encode()).hexdigest(),text(row["Answer Hash"]) or hashlib.sha256(f"answer:{row['Session ID']}:{row['Position']}".encode()).hexdigest(),text(row["Submit Status"]) or "draft",f"legacy:{row['Session ID']}:{row['Position']}:{integer(row['Revision'],1)}",text(row["Contract Version"]),timestamp_value(row["Updated At"]) or import_timestamp,timestamp_value(row["Updated At"]) or import_timestamp)
            for row in sheets["Answer Drafts"]
        ))
        emit_copy(out, "english_private.answer_draft_history", ["id","owner_id","legacy_history_id","answer_draft_id","session_id","question_id","position","previous_answer","previous_revision","next_answer","next_revision","event_type","client_instance_id","page_started_at","answer_hash","contract_version","recorded_at"], (
            (stable_id("draft-history",row["History ID"]),owner,text(row["History ID"]),draft_ids.get((text(row["Session ID"]),integer(row["Position"]))),sessions[text(row["Session ID"])],session_questions.get((text(row["Session ID"]),integer(row["Position"]))),integer(row["Position"]),text(row["Previous Answer"]),integer(row["Previous Revision"]),text(row["Next Answer"]),integer(row["Next Revision"],1),text(row["Event Type"]) or "legacy",text(row["Client Instance ID"]),timestamp_value(row["Page Started At"]),text(row["Answer Hash"]),text(row["Contract Version"]),timestamp_value(row["Recorded At"]) or import_timestamp)
            for row in sheets["Answer Draft History"]
        ))

        journal_by_submission = {text(row["Submission ID"]): row for row in sheets["Commit Journal"]}
        emit_copy(out, "english_private.submissions", ["id","owner_id","legacy_submission_id","session_id","queue_id","revision","frozen_hash","idempotency_key","status","error_code","error_detail","created_at","updated_at","completed_at"], (
            (submissions[key],owner,key,sessions[text(row["Session ID"])],queues[text(row["Queue ID"])],1,text(row["Answer Hash"]) or hashlib.sha256(key.encode()).hexdigest(),f"legacy:{key}","committed" if text(row["Status"]) == "committed" else "failed",text(row["Error Code"]),text(row["Error Detail"]),timestamp_value(row["Started At"]) or import_timestamp,timestamp_value(row["Updated At"]) or import_timestamp,timestamp_value(row["Completed At"]))
            for key,row in journal_by_submission.items()
        ))
        emit_copy(out, "english_private.grade_requests", ["id","owner_id","submission_id","question_id","position","phrase_id","candidate_id","observed_answer","prompt_zh","prompt_en","expected_answers","accepted_variants","semantic_boundary","grading_rubric","review_stage","answer_hash","request_status","snapshot_contract_version","contract_version","created_at"], (
            (request_ids[(text(row["Submission ID"]),integer(row["Position"]))],owner,submissions[text(row["Submission ID"])],session_questions[(text(row["Session ID"]),integer(row["Position"]))],integer(row["Position"]),phrases.get(text(row["Phrase ID"])),candidates.get(text(row["Candidate ID"])),text(row["Observed Answer"]) or "[UNKNOWN]",text(row["Prompt ZH"]),text(row["Prompt EN"]),json_data(row["Expected Answers JSON"],[]),json_data(row["Accepted Variants JSON"],[]),text(row["Semantic Boundary"]),text(row["Grading Rubric"]),integer(row["Review Stage"]),text(row["Answer Hash"]) or hashlib.sha256(f"{row['Submission ID']}:{row['Position']}".encode()).hexdigest(),"legacy",text(row["Snapshot Contract Version"]) or "1.0",text(row["Contract Version"]),timestamp_value(row["Created At"]) or import_timestamp)
            for row in sheets["Grade Requests"]
        ))
        emit_copy(out, "english_private.grade_result_attempts", ["id","owner_id","submission_id","legacy_submission_id","position","phrase_id","candidate_id","answer_hash","result","feedback_zh","error_category","confidence","evidence","expected_answer","observed_answer","candidate_suggestions","extra_practice","grading_batch_id","prompt_version","grade_status","contract_version","created_at"], (
            (stable_id("grade-attempt",f"{row['__source_row']}:{row['__row_hash']}"),owner,submissions.get(text(row["Submission ID"])),text(row["Submission ID"]),integer(row["Position"]),phrases.get(text(row["Phrase ID"])),candidates.get(text(row["Candidate ID"])),text(row["Answer Hash"]),text(row["Result"]),text(row["Feedback ZH"]),text(row["Error Category"]),number(row["Confidence"]),text(row["Evidence"]),text(row["Expected Answer"]),text(row["Observed Answer"]),json_data(row["Candidate Suggestions JSON"],[]),json_data(row["Extra Practice JSON"],[]),text(row["Grading Batch ID"]) or f"legacy-row-{row['__source_row']}",text(row["Prompt Version"]),text(row["Grade Status"]) or "legacy",text(row["Contract Version"]),timestamp_value(row["Created At"]))
            for row in sheets["Grade Inbox"]
        ))
        emit_copy(out, "english_private.commit_journal", ["id","owner_id","submission_id","answer_hash","status","last_completed_step","error_code","error_detail","readback_status","result","confirmation","contract_version","started_at","updated_at","completed_at"], (
            (stable_id("commit-journal",row["Submission ID"]),owner,submissions[text(row["Submission ID"])],text(row["Answer Hash"]) or hashlib.sha256(str(row["Submission ID"]).encode()).hexdigest(),text(row["Status"]) or "committed",text(row["Last Completed Step"]),text(row["Error Code"]),text(row["Error Detail"]),text(row["Readback Status"]),json_data(row["Result JSON"],{}),json_data(row["Confirmation JSON"],{}),text(row["Contract Version"]),timestamp_value(row["Started At"]) or import_timestamp,timestamp_value(row["Updated At"]) or import_timestamp,timestamp_value(row["Completed At"]))
            for row in sheets["Commit Journal"]
        ))

        emit_copy(out, "english_private.review_events", ["id","owner_id","phrase_id","session_id","question_position","prompt","expected_answer","user_answer","result","tag","follow_up_needed","notes","legacy_attempt_id","attempt_type","parent_attempt_id","question_type","affects_srs","contract_version","reviewed_at"], (
            (stable_id("review-event",f"{row['__source_row']}:{row['__row_hash']}"),owner,phrase_for_review(row,session_questions,sheets,phrases),sessions.get(text(row["Session ID"])),integer(row["Question #"]),text(row["Prompt"]),text(row["Expected Answer"]),text(row["User Answer"]),text(row["Result"]),text(row["Tag"]),boolean(row["Follow-up Needed"]),text(row["Notes"]),text(row["Attempt ID"]),text(row["Attempt Type"]),text(row["Parent Attempt ID"]),text(row["Question Type"]),boolean(row["Affects SRS?"]) is True,text(row["Contract Version"]),timestamp_value(row["Date"]) or import_timestamp)
            for row in sheets["Review Log"]
        ))
        emit_copy(out, "english_private.error_events", ["id","owner_id","legacy_error_id","phrase_id","session_id","chunk","error_type","user_answer","correction","explanation","next_action","resolved","attempt_id","occurred_at","resolved_at","contract_version","created_at"], (
            (stable_id("error-event",row.get("Error ID") or row["__row_hash"]),owner,text(row["Error ID"]),phrases.get(text(row["Phrase ID"])),sessions.get(text(row["Session ID"])),text(row["Chunk"]),text(row["Error Type"]),text(row["User Answer"]),text(row["Correction"]),text(row["Explanation"]),text(row["Next Action"]),boolean(row["Resolved?"]),text(row["Attempt ID"]),timestamp_value(row["Date"]),timestamp_value(row["Resolution Date"]),text(row["Contract Version"]),timestamp_value(row["Date"]) or import_timestamp)
            for row in sheets["Error Log"]
        ))
        decision_status = {"committed":"committed","explanation_only":"rejected"}
        emit_copy(out, "english_private.context_candidates", ["id","owner_id","context_id","position","selected_text","candidate","cue_zh","candidate_type","context_meaning","why_useful","topic","difficulty","natural_example","common_mistake","extraction_rationale","confidence","decision_status","edited_candidate","processing_batch_id","candidate_id","contract_version","created_at","committed_at"], (
            (stable_id("context-candidate",f"{row['Context ID']}:{row['Proposal Position']}"),owner,contexts[text(row["Context ID"])],integer(row["Proposal Position"]),text(row["Selected Text"]),text(row["Candidate"]) or "[UNKNOWN]",text(row["Chinese Cue"]),text(row["Candidate Type"]),text(row["Context Meaning"]),text(row["Why Useful"]),text(row["Topic"]),text(row["Difficulty"]),text(row["Natural Example"]),text(row["Common Mistake"]),text(row["Extraction Rationale"]),number(row["Confidence"]),decision_status.get(text(row["Decision Status"]),"pending"),text(row["Edited Candidate"]),text(row["Processing Batch ID"]),candidates.get(text(row["Candidate ID"])),text(row["Contract Version"]),timestamp_value(row["Created At"]) or import_timestamp,timestamp_value(row["Committed At"]))
            for row in sheets["Context Candidate Inbox"]
        ))

        for session_key in sorted(draft_sessions):
            rows = [row for row in sheets["Answer Drafts"] if text(row["Session ID"]) == session_key]
            payload = json.dumps([{key:value for key,value in row.items() if not key.startswith("__")} for row in rows],ensure_ascii=False,separators=(",",":"))
            out.write(
                "insert into english_private.legacy_recovery(id,owner_id,recovery_kind,legacy_key,payload,learning_date,visible,automatically_resumable) values "
                f"('{stable_id('legacy-recovery',session_key)}','{owner}','answer_draft_session',{sql_literal(session_key)},{sql_literal(payload)}::jsonb,{sql_literal(session_key[:10])}::date,true,false);\n"
            )
        out.write(
            "update english_private.import_rows set promotion_status='promoted' where import_batch_id="
            f"'{batch_id}' and promotion_status='staged';\n"
        )
        report = {
            "sourceSheets": manifest["sheet_count"], "sourceRows": sum(s["row_count"] for s in manifest["sheets"]),
            "sourceFormulas": manifest["formula_count"], "sourceSnapshotHash": manifest["snapshot_sha256"],
            "syntheticLegacyQueues": sum(1 for key in queue_groups if key.startswith("legacy-session:")),
        }
        out.write(
            "update english_private.import_batches set status='promoted',reconciliation="
            f"{sql_literal(json.dumps(report,separators=(',',':')))}::jsonb,reconciled_at=now(),promoted_at=now() where id='{batch_id}';\n"
        )
        out.write("commit;\n")

    tables = {**stage.tables, **promote.tables}
    recovery_rows = [{"id": stable_id("legacy-recovery", key), "owner_id": owner} for key in sorted(draft_sessions)]
    tables["english_private.legacy_recovery"] = recovery_rows
    target_audit = validate_targets(tables)
    report.update({"version": 2, "ownerId": owner, "batchId": batch_id,
                   "sourceState": manifest.get("source_state", "rehearsal"),
                   "sourceFrozenAt": manifest.get("source_frozen_at"),
                   "importedAt": import_timestamp, "sourceAudit": source_audit, "syntheticSessionCount": len(session_records) - len(sheets["Session Log"]),
                   "sheetCounts": {sheet["sheet"]: sheet["row_count"] for sheet in manifest["sheets"]},
                   "tableCounts": {table: len(rows) for table, rows in tables.items()},
                   "targetIds": {table: sorted(row["id"] for row in rows) for table, rows in tables.items()},
                   "rawRows": [{"sourceSheet": row["source_sheet"], "sourceRow": row["source_row"],
                                "rowHash": row["row_hash"], "rawHash": canonical_hash(row["raw_json"]),
                                "formulaHash": canonical_hash(row["formula_cells"]),
                                "targetTable": row["target_table"], "targetId": row["target_id"]}
                               for row in tables["english_private.import_rows"]],
                   "targetAudit": target_audit,
                   "legacyPolicy": "Queues and sessions are historical; drafts never automatically resume. Missing created_at uses importedAt; source cells remain authoritative."})
    # Native psql performs the full import in one transaction.
    start = stage.getvalue().removeprefix("begin;\n").removesuffix("commit;\n")
    finish = promote.getvalue().removeprefix("begin;\n").removesuffix("commit;\n")
    preflight = sql_preflight(owner, batch_id, tables)
    server_report = {key: report[key] for key in ("version", "sourceSheets", "sourceRows", "sourceFormulas", "sourceSnapshotHash", "sheetCounts", "tableCounts", "syntheticSessionCount", "legacyPolicy")}
    finish += "update english_private.import_batches set reconciliation=" + sql_literal(json.dumps(server_report,ensure_ascii=False)) + f"::jsonb where id='{batch_id}' and owner_id='{owner}';\n"
    assertions = sql_assertions(owner, batch_id, manifest, tables)
    sql = "\\set ON_ERROR_STOP on\nbegin;\n" + preflight + start + finish + assertions + "commit;\n"
    sql_path = output / "import.sql"
    output.mkdir(parents=True, exist_ok=True)
    sql_path.write_text(sql, encoding="utf-8")
    sql_path.chmod(0o600)
    report["sqlSha256"] = hashlib.sha256(sql.encode()).hexdigest()
    api_sql = sql.removeprefix("\\set ON_ERROR_STOP on\n")
    for copy_sql, insert_sql in stage.copy_blocks + promote.copy_blocks:
        api_sql = api_sql.replace(copy_sql, insert_sql, 1)
    # db query can use prepared statements: submit one top-level DO statement.
    # DO is atomic without any explicit transaction commands inside its body.
    api_body = api_sql.removeprefix("begin;\n").removesuffix("commit;\n")
    api_body = api_body.replace("select pg_advisory_xact_lock(", "perform pg_advisory_xact_lock(")
    api_body = api_body.replace("\nselect set_config(", "\nperform set_config(")
    delimiter = "$import_" + hashlib.sha256(api_body.encode()).hexdigest()[:16] + "$"
    if delimiter in api_body:
        raise ValueError("Unexpected SQL delimiter collision")
    api_sql = f"do {delimiter} begin\n" + api_body + f"end {delimiter};\n"
    api_path = output / "import-api.sql"
    api_path.write_text(api_sql, encoding="utf-8")
    api_path.chmod(0o600)
    report["apiSqlSha256"] = hashlib.sha256(api_sql.encode()).hexdigest()
    from staged_import import build_staged_files
    report["staged"] = build_staged_files(output, owner, batch_id, manifest, report, start + finish,
                                          stage, promote, assertions)
    write_json(output / "import-manifest.json", report)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("snapshot", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--owner-id", required=True, type=uuid.UUID)
    parser.add_argument("--imported-at", help="Optional fixed ISO timestamp for deterministic rehearsal output")
    args = parser.parse_args()
    report = build_import(args.snapshot, args.output, str(args.owner_id), imported_at=args.imported_at)
    print(json.dumps({key: report[key] for key in ("batchId", "sourceRows", "sourceFormulas", "sourceSnapshotHash", "tableCounts")}, ensure_ascii=False))


def phrase_for_review(row, session_questions, sheets, phrases):
    tagged = phrases.get(text(row.get("Tag")))
    if tagged:
        return tagged
    session_key = text(row.get("Session ID"))
    position = integer(row.get("Question #"))
    if session_key and position:
        for question in sheets["Session Questions"]:
            if text(question.get("Session ID")) == session_key and integer(question.get("Position")) == position and text(question.get("Question Status")) != "rejected":
                return phrases.get(text(question.get("Phrase ID")))
    expected = text(row.get("Expected Answer"))
    if expected:
        for phrase in sheets["Phrase Bank"]:
            if text(phrase.get("Chunk")) == expected:
                return phrases.get(text(phrase.get("ID")))
    return None


def sql_literal(value: str | None) -> str:
    if value is None:
        return "null"
    return "'" + value.replace("'", "''") + "'"


if __name__ == "__main__":
    main()
