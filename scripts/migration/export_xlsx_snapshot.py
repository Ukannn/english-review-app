#!/usr/bin/env python3
"""Export all populated XLSX cells, including formula-only rows, without changing it."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from datetime import date, datetime, time, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("Non-finite spreadsheet value")
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    raise ValueError(f"Unsupported spreadsheet value type: {type(value).__name__}")


def canonical_hash(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":"), allow_nan=False).encode()).hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    path.chmod(0o600)


def validate_freeze_metadata(source_state, frozen_at, exported_at):
    if source_state not in ("rehearsal", "final_frozen"):
        raise ValueError("Unknown source state")
    if source_state == "rehearsal":
        if frozen_at:
            raise ValueError("A rehearsal snapshot cannot claim a freeze timestamp")
        return
    values = []
    for value in (frozen_at, exported_at):
        if not isinstance(value, str):
            raise ValueError("Final snapshot requires timezone-aware freeze and export timestamps")
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            raise ValueError("Final snapshot requires timezone-aware freeze and export timestamps")
        values.append(parsed)
    if values[0] > values[1]:
        raise ValueError("The legacy write freeze must precede the final snapshot export")


def export_snapshot(xlsx: Path, output: Path, *, source_state="rehearsal", frozen_at=None) -> dict:
    exported_at = datetime.now(timezone.utc).isoformat()
    validate_freeze_metadata(source_state, frozen_at, exported_at)
    if output.exists() and any(output.iterdir()):
        raise ValueError("Snapshot output must be empty; existing evidence is never overwritten")
    source_hash = hashlib.sha256(xlsx.read_bytes()).hexdigest()
    formula_book = load_workbook(xlsx, data_only=False)
    value_book = load_workbook(xlsx, data_only=True)
    manifest = {
        "format_version": 2, "source": xlsx.name, "source_sha256": source_hash,
        "source_modified_at": datetime.fromtimestamp(xlsx.stat().st_mtime, timezone.utc).isoformat(),
        "exported_at": exported_at,
        "excel_epoch": formula_book.epoch.isoformat(), "timezone": "Asia/Shanghai",
        "source_state": source_state, "source_frozen_at": frozen_at,
        "sheet_count": len(formula_book.sheetnames), "sheets": [], "formula_count": 0,
    }
    try:
        for index, name in enumerate(formula_book.sheetnames, 1):
            sheet, cached = formula_book[name], value_book[name]
            populated = [(cell.row, cell.column) for row in sheet.iter_rows() for cell in row if cell.value is not None]
            last_row = max((r for r, _ in populated), default=0)
            last_col = max((c for _, c in populated), default=0)
            headers = [json_value(sheet.cell(1, c).value) for c in range(1, last_col + 1)] if last_row else []
            rows, formulas = [], []
            for r in range(2, last_row + 1):
                if not any(sheet.cell(r, c).value is not None for c in range(1, last_col + 1)):
                    continue
                values = [json_value(cached.cell(r, c).value) for c in range(1, last_col + 1)]
                row_formulas = []
                for c in range(1, last_col + 1):
                    cell = sheet.cell(r, c)
                    if cell.data_type == "f":
                        row_formulas.append({"cell": cell.coordinate, "source_row": r,
                                             "header": headers[c-1], "formula": cell.value,
                                             "calculated_value": values[c-1]})
                record = {"source_sheet": name, "source_row": r, "values": values,
                          "formula_cells": row_formulas}
                record["row_hash"] = canonical_hash(record)
                rows.append(record)
                formulas.extend(row_formulas)
            payload = {"sheet": name, "headers": headers, "row_count": len(rows),
                       "formula_count": len(formulas), "rows": rows, "formulas": formulas}
            # Numbered names avoid path traversal and collisions from sanitized sheet names.
            filename = f"sheets/{index:03d}.json"
            write_json(output / filename, payload)
            manifest["sheets"].append({"sheet": name, "file": filename, "headers": headers,
                                       "row_count": len(rows), "formula_count": len(formulas),
                                       "content_sha256": canonical_hash(payload)})
            manifest["formula_count"] += len(formulas)
        manifest["snapshot_sha256"] = canonical_hash(manifest["sheets"])
        manifest["manifest_sha256"] = canonical_hash(manifest)
        if hashlib.sha256(xlsx.read_bytes()).hexdigest() != source_hash:
            raise ValueError("Source workbook changed during export; discard this snapshot")
        write_json(output / "manifest.json", manifest)
        return manifest
    finally:
        formula_book.close()
        value_book.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--source-state", choices=("rehearsal", "final_frozen"), default="rehearsal")
    parser.add_argument("--frozen-at", help="Operator-confirmed legacy write freeze time; required for final_frozen")
    args = parser.parse_args()
    manifest = export_snapshot(args.xlsx, args.output, source_state=args.source_state, frozen_at=args.frozen_at)
    print(json.dumps({"sheetCount": manifest["sheet_count"],
                      "rowCount": sum(s["row_count"] for s in manifest["sheets"]),
                      "formulaCount": manifest["formula_count"],
                      "snapshotHash": manifest["snapshot_sha256"]}))


if __name__ == "__main__":
    main()
